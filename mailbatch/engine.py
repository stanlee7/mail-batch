# -*- coding: utf-8 -*-
"""메일 배치 코어 엔진.

엑셀 명단을 읽어 {{컬럼명}} 자리표시를 채운 개인별 메일을 SMTP로 발송한다.
GUI 없이도 동작하며, 네트워크가 필요한 부분은 send_batch뿐이다.
"""
import csv
import glob as _glob
import mimetypes
import os
import re
import smtplib
import ssl
import time
from dataclasses import dataclass, field
from email.message import EmailMessage
from email.utils import formataddr
from typing import Callable, Optional

import openpyxl

PLACEHOLDER = re.compile(r"\{\{([^{}]+)\}\}")

ProgressCb = Optional[Callable[[int, int, str], None]]
CancelCb = Optional[Callable[[], bool]]

_ILLEGAL_FILENAME = re.compile(r'[\\/:*?"<>|\r\n\t]')
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

# 연속 실패가 이만큼 쌓이면 일일 한도/차단으로 보고 중단한다.
MAX_CONSECUTIVE_FAILURES = 5

ATTACH_WARN_BYTES = 10 * 1024 * 1024  # 메일당 첨부 합계 경고 기준 (네이버/다음 한도권)


def sanitize_filename(name: str) -> str:
    name = _ILLEGAL_FILENAME.sub("_", str(name)).strip().rstrip(".")
    return name[:120] or "문서"


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_excel(path: str):
    """첫 행을 헤더로 읽어 (headers, rows[dict]) 반환. 빈 행은 건너뛴다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        wb.close()
        return [], []
    headers = [str(h).strip() for h in header_row if h is not None and str(h).strip()]
    n = len(headers)
    rows = []
    for raw in rows_iter:
        if raw is None or all(v is None or str(v).strip() == "" for v in raw[:n]):
            continue
        row = {}
        for i, h in enumerate(headers):
            v = raw[i] if i < len(raw) else None
            row[h] = _cell_to_str(v)
        rows.append(row)
    wb.close()
    return headers, rows


def render_template(text: str, row: dict) -> str:
    """'{{컬럼명}}' 자리표시를 행 데이터로 채운다."""
    return PLACEHOLDER.sub(lambda m: row.get(m.group(1).strip(), ""), text or "")


def find_unknown_placeholders(text: str, headers: list) -> list:
    """텍스트에 있지만 엑셀 헤더에 없는 자리표시 목록 (발송 전 오타 검증)."""
    known = {h.strip() for h in headers}
    seen = []
    for m in PLACEHOLDER.finditer(text or ""):
        name = m.group(1).strip()
        if name not in known and name not in seen:
            seen.append(name)
    return seen


def validate_email(addr: str) -> bool:
    return bool(_EMAIL_RE.match((addr or "").strip()))


def guess_email_column(headers: list, rows: list) -> str:
    """헤더 이름 또는 데이터 내용으로 이메일 컬럼을 추정한다."""
    for h in headers:
        low = h.lower()
        if "메일" in h or "email" in low or "e-mail" in low:
            return h
    for h in headers:
        sample = [r.get(h, "") for r in rows[:10] if r.get(h, "")]
        if sample and sum(1 for v in sample if "@" in v) >= len(sample) * 0.8:
            return h
    return headers[0] if headers else ""


def match_attachments(pattern: str, folder: str, row: dict):
    """개인별 첨부 패턴('{이름}_수료증.pdf', '*' 허용)을 행 데이터로 채워 폴더에서 찾는다.

    반환: (paths, error)  — 못 찾으면 paths=[], error=메시지
    """
    name = re.sub(r"\{([^{}]+)\}", lambda m: row.get(m.group(1).strip(), ""), pattern or "").strip()
    if not name:
        return [], "첨부 패턴이 비어 있습니다"
    if "*" in name or "?" in name:
        paths = sorted(_glob.glob(os.path.join(folder, name)))
        if not paths:
            return [], f"'{name}' 패턴과 일치하는 파일이 없습니다"
        return [os.path.abspath(p) for p in paths], None
    path = os.path.join(folder, name)
    if not os.path.exists(path):
        return [], f"'{name}' 파일이 없습니다"
    return [os.path.abspath(path)], None


@dataclass
class SendConfig:
    host: str = ""
    port: int = 465
    security: str = "ssl"            # "ssl" | "starttls" | "none"(테스트용)
    user: str = ""
    password: str = ""
    sender_name: str = ""
    to_column: str = ""
    subject_tpl: str = ""
    body_tpl: str = ""
    body_is_html: bool = False
    fixed_attachments: list = field(default_factory=list)
    attach_dir: str = ""             # 개인별 첨부 폴더 (빈 값이면 미사용)
    attach_pattern: str = ""
    missing_attach_policy: str = "fail"   # "fail" | "skip_row" | "send_without"
    delay_sec: float = 2.0


class AuthError(Exception):
    """로그인 실패 등 발송 전체를 중단해야 하는 오류."""


def _strip_header(value: str) -> str:
    """제목/주소 헤더에 줄바꿈이 끼어드는 것(헤더 인젝션)을 막는다."""
    return re.sub(r"[\r\n]+", " ", value or "").strip()


def _row_attachments(cfg: SendConfig, row: dict):
    """행 하나의 첨부 목록을 계산한다. 반환: (paths, missing_error)"""
    paths = list(cfg.fixed_attachments)
    err = None
    if cfg.attach_dir and cfg.attach_pattern:
        matched, err = match_attachments(cfg.attach_pattern, cfg.attach_dir, row)
        paths += matched
    return paths, err


def build_message(cfg: SendConfig, row: dict, to_override: str = None) -> EmailMessage:
    """행 하나로 EmailMessage를 만든다. 첨부 누락 정책은 호출자가 처리."""
    msg = EmailMessage()
    sender = formataddr((cfg.sender_name, cfg.user)) if cfg.sender_name else cfg.user
    msg["From"] = sender
    msg["To"] = _strip_header(to_override or row.get(cfg.to_column, ""))
    msg["Subject"] = _strip_header(render_template(cfg.subject_tpl, row))
    body = render_template(cfg.body_tpl, row)
    if cfg.body_is_html:
        msg.set_content("HTML 메일입니다. HTML을 지원하는 메일 클라이언트로 확인하세요.")
        msg.add_alternative(body, subtype="html")
    else:
        msg.set_content(body)

    paths, _ = _row_attachments(cfg, row)
    for path in paths:
        ctype, _enc = mimetypes.guess_type(path)
        maintype, _, subtype = (ctype or "application/octet-stream").partition("/")
        with open(path, "rb") as fp:
            msg.add_attachment(fp.read(), maintype=maintype, subtype=subtype,
                               filename=os.path.basename(path))
    return msg


def dry_run_row(cfg: SendConfig, row: dict, index: int) -> dict:
    """발송 없이 행 하나의 렌더 결과를 보여준다 (미리보기용)."""
    paths, attach_err = _row_attachments(cfg, row)
    total = sum(os.path.getsize(p) for p in paths if os.path.exists(p))
    warnings = []
    if attach_err:
        warnings.append(f"개인별 첨부 누락: {attach_err}")
    email = row.get(cfg.to_column, "")
    if not validate_email(email):
        warnings.append(f"이메일 주소가 올바르지 않습니다: '{email}'")
    if total > ATTACH_WARN_BYTES:
        warnings.append(f"첨부 합계 {total / 1024 / 1024:.1f}MB — 수신 측 한도를 넘을 수 있습니다")
    return {
        "index": index,
        "email": email,
        "subject": _strip_header(render_template(cfg.subject_tpl, row)),
        "body": render_template(cfg.body_tpl, row),
        "attachments": paths,
        "warnings": warnings,
    }


class _Sender:
    """SMTP 연결 1개를 감싼다. 끊기면 1회 재접속한다."""

    def __init__(self, cfg: SendConfig):
        self.cfg = cfg
        self.smtp = None

    def connect(self):
        cfg = self.cfg
        if cfg.security == "ssl":
            self.smtp = smtplib.SMTP_SSL(cfg.host, cfg.port, timeout=30,
                                         context=ssl.create_default_context())
        else:
            self.smtp = smtplib.SMTP(cfg.host, cfg.port, timeout=30)
            if cfg.security == "starttls":
                self.smtp.starttls(context=ssl.create_default_context())
        if cfg.user and cfg.password:
            try:
                self.smtp.login(cfg.user, cfg.password)
            except smtplib.SMTPAuthenticationError as e:
                raise AuthError(
                    f"로그인 실패 ({e.smtp_code}): 아이디/앱 비밀번호를 확인하세요. "
                    "네이버·지메일은 일반 비밀번호가 아니라 '앱 비밀번호'가 필요합니다.") from e

    def send(self, msg: EmailMessage):
        if self.smtp is None:
            self.connect()
        try:
            self.smtp.send_message(msg)
        except (smtplib.SMTPServerDisconnected, ConnectionError, OSError):
            self.close()
            self.connect()
            self.smtp.send_message(msg)

    def close(self):
        if self.smtp is not None:
            try:
                self.smtp.quit()
            except Exception:
                pass
            self.smtp = None


def _sleep_cancellable(seconds: float, cancel: CancelCb) -> bool:
    """간격 대기. 중간에 취소되면 True 반환."""
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        if cancel and cancel():
            return True
        time.sleep(min(0.2, max(0.0, end - time.monotonic())))
    return bool(cancel and cancel())


def send_batch(cfg: SendConfig, rows: list, progress: ProgressCb = None,
               cancel: CancelCb = None, to_override: str = None) -> list:
    """행마다 개인화 메일을 발송한다.

    반환: [{"index", "email", "subject", "attachments", "error": None|str}, ...]
    - 로그인 실패는 AuthError로 전체 중단
    - 행별 오류는 기록 후 계속, 연속 MAX_CONSECUTIVE_FAILURES건이면 중단
    """
    results = []
    total = len(rows)
    consecutive = 0
    sender = _Sender(cfg)
    try:
        sender.connect()  # 자격 증명 문제는 첫 행 전에 드러낸다
        for i, row in enumerate(rows):
            if cancel and cancel():
                break
            email = _strip_header(to_override or row.get(cfg.to_column, ""))
            item = {"index": i, "email": email,
                    "subject": _strip_header(render_template(cfg.subject_tpl, row)),
                    "attachments": [], "error": None}
            paths, attach_err = _row_attachments(cfg, row)
            item["attachments"] = paths
            try:
                if not validate_email(email):
                    raise ValueError(f"이메일 주소 오류: '{email}'")
                if attach_err and cfg.missing_attach_policy == "fail":
                    raise FileNotFoundError(attach_err)
                if attach_err and cfg.missing_attach_policy == "skip_row":
                    item["error"] = f"건너뜀 — {attach_err}"
                    results.append(item)
                    if progress:
                        progress(i + 1, total, f"{email} (건너뜀)")
                    continue
                msg = build_message(cfg, row, to_override=to_override)
                sender.send(msg)
                consecutive = 0
            except AuthError:
                raise
            except Exception as e:
                item["error"] = str(e)
                consecutive += 1
            results.append(item)
            if progress:
                status = item["error"] or "발송"
                progress(i + 1, total, f"{email} — {status}")
            if consecutive >= MAX_CONSECUTIVE_FAILURES:
                raise AuthError(
                    f"연속 {consecutive}건 실패로 중단했습니다. "
                    "일일 발송 한도 초과 또는 차단일 수 있습니다. 내일 다시 시도하거나 발송 간격을 늘려보세요.")
            if i < total - 1 and cfg.delay_sec > 0:
                if _sleep_cancellable(cfg.delay_sec, cancel):
                    break
    finally:
        sender.close()
    return results


def export_results(results: list, path: str):
    """발송 결과를 xlsx 또는 csv로 저장한다."""
    headers = ["행", "이메일", "제목", "첨부", "결과"]
    rows = [[r["index"] + 1, r["email"], r["subject"],
             "; ".join(os.path.basename(p) for p in r["attachments"]),
             r["error"] or "성공"] for r in results]
    if path.lower().endswith(".csv"):
        with open(path, "w", newline="", encoding="utf-8-sig") as fp:
            w = csv.writer(fp)
            w.writerow(headers)
            w.writerows(rows)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "발송 결과"
        ws.append(headers)
        for r in rows:
            ws.append(r)
        wb.save(path)
